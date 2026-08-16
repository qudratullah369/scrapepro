from openpyxl import load_workbook

from scrapepro.core.record import Record
from scrapepro.exporters.excel import ExcelExporter


def test_excel_exporter_writes_header_and_records(tmp_path):
    records = [
        Record(
            name="Cafe A",
            category="cafe",
            phone="+923001234567",
            website="https://example.com",
            address="Main Street, Islamabad, Pakistan",
            city="Islamabad",
            country="Pakistan",
            latitude=33.6844,
            longitude=73.0479,
            rating=4.7,
            reviews=120,
            source="google_maps",
            place_id="p1",
        )
    ]

    output_path = tmp_path / "cafes.xlsx"

    result = ExcelExporter().export(records, output_path)

    assert result == output_path
    assert output_path.exists()

    workbook = load_workbook(output_path)

    worksheet = workbook["Records"]

    assert worksheet.max_row == 2
    assert worksheet.max_column == len(ExcelExporter.fieldnames)

    assert worksheet["A1"].value == "name"
    assert worksheet["A2"].value == "Cafe A"

    assert worksheet["F2"].value == "Islamabad"
    assert worksheet["G2"].value == "Pakistan"

    assert worksheet["J2"].value == 4.7
    assert worksheet["K2"].value == 120
    assert worksheet["L2"].value == "google_maps"
    assert worksheet["M2"].value == "p1"

    workbook.close()


def test_excel_exporter_writes_empty_dataset_with_header(tmp_path):
    output_path = tmp_path / "empty.xlsx"

    ExcelExporter().export([], output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Records"]

    assert worksheet.max_row == 1
    assert [
        cell.value
        for cell in worksheet[1]
    ] == ExcelExporter.fieldnames

    workbook.close()


def test_excel_exporter_preserves_text(tmp_path):
    records = [
        Record(
            name="Tea House",
            address="Islamabad, Pakistan",
            city="Islamabad",
            country="Pakistan",
        )
    ]

    output_path = tmp_path / "text.xlsx"

    ExcelExporter().export(records, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Records"]

    assert worksheet["A2"].value == "Tea House"
    assert worksheet["E2"].value == "Islamabad, Pakistan"
    assert worksheet["F2"].value == "Islamabad"
    assert worksheet["G2"].value == "Pakistan"

    workbook.close()
